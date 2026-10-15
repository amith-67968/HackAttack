import asyncio
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.api.research import get_repository
from app.core.config import Settings, get_settings
from app.main import app
from app.repositories.tasks import SupabaseTaskRepository
from app.services.artifacts import build_artifact
from app.services.mcp_search import discover_and_search
from app.services.retry import classify_error


def evidence_fixture():
    task_id = str(uuid4())
    task = {"id": task_id, "query": "Compare Alpha clinic fees", "status": "COMPLETED_WITH_LIMITATIONS",
            "plan": {"required_fields": ["price", "eligibility"]}}
    sources = [{"id": "source-1", "url": "https://example.org/fees", "title": "Fees", "page_content": "Alpha Clinic consultation price is $50. Eligibility: adults."}]
    findings = [
        {"source_id": "source-1", "entity_name": "Alpha Clinic", "field_name": "price", "field_value": "$50", "value_status": "found", "verification_status": "source_supported", "supporting_excerpt": "consultation price is $50"},
        {"source_id": "source-1", "entity_name": "Alpha Clinic", "field_name": "eligibility", "field_value": "adults", "value_status": "found", "verification_status": "source_supported", "supporting_excerpt": "Eligibility: adults"},
        {"source_id": "source-1", "entity_name": "Alpha Clinic", "field_name": "fraudulent claim", "field_value": "free", "value_status": "found", "verification_status": "source_supported", "supporting_excerpt": "Not in source"},
    ]
    return task, sources, findings


def test_spreadsheet_has_three_reopened_sheets_and_excludes_unsupported_values():
    task, sources, findings = evidence_fixture()
    sources[0]["page_content"] += " Special offer =1+1."
    findings.append({"source_id": "source-1", "entity_name": "Alpha Clinic", "field_name": "offer", "field_value": "=1+1", "value_status": "found", "verification_status": "source_supported", "supporting_excerpt": "Special offer =1+1."})
    payload, preview = build_artifact(task, findings, sources, "spreadsheet")
    book = load_workbook(__import__("io").BytesIO(payload), read_only=True)
    assert book.sheetnames == ["Results", "Sources & Evidence", "Research Summary"]
    assert preview["task_status"] == "COMPLETED_WITH_LIMITATIONS"
    assert all(row["field"] != "fraudulent claim" for row in preview["rows"])
    assert any(row["field"] == "eligibility" and row["status"] == "source-supported" for row in preview["rows"])
    assert book["Results"]["C3"].value == "'=1+1"


def test_email_and_report_are_honest_source_backed_previews():
    task, sources, findings = evidence_fixture()
    email, preview = build_artifact(task, findings, sources, "email")
    assert b"has not been sent" in email and preview["draft_only"] is True
    report, report_preview = build_artifact(task, findings, sources, "report")
    assert b"COMPLETED_WITH_LIMITATIONS" in report
    assert "https://example.org/fees" in report_preview["markdown"]


@pytest.mark.parametrize("status", ["RUNNING", "FAILED", "COMPLETED_WITH_LIMITATIONS"])
def test_incomplete_and_failed_artifacts_preserve_truthful_task_status(status):
    task, sources, findings = evidence_fixture()
    task["status"] = status
    for artifact_type in ("spreadsheet", "report", "email"):
        payload, preview = build_artifact(task, findings, sources, artifact_type)
        assert status in str(preview)
        if artifact_type == "spreadsheet":
            assert load_workbook(__import__("io").BytesIO(payload), read_only=True)["Research Summary"]["B2"].value == status
        else:
            assert status.encode() in payload
        assert task["status"] == status


def test_mcp_protocol_discovers_and_calls_search_tool():
    class Search:
        async def search(self, query, *, domain, max_results):
            assert query == "focused question" and max_results == 3
            return [{"url": "https://example.org", "title": "Example"}]

    results = asyncio.run(discover_and_search(Search(), "focused question", domain=None, max_results=3))
    assert results == [{"url": "https://example.org", "title": "Example"}]


def test_mcp_provider_failure_keeps_retry_classification_without_raw_error():
    class Search:
        async def search(self, *_args, **_kwargs):
            error = RuntimeError("sensitive provider response and key")
            error.status_code = 503
            raise error

    try:
        asyncio.run(discover_and_search(Search(), "query", domain=None, max_results=2))
        assert False, "provider failure expected"
    except Exception as exc:
        classification = classify_error(exc)
        assert classification.name == "temporary_server_error" and classification.retryable
        assert "sensitive" not in str(exc)


def test_supabase_artifact_storage_upload_and_signed_url_use_private_backend_client():
    class Bucket:
        def upload(self, path, payload, options):
            assert path.endswith(".xlsx") and payload == b"binary"
            assert options["content-type"].startswith("application/") and options["upsert"] == "true"
        def create_signed_url(self, path, expiry):
            assert path.startswith("task/") and expiry == 300
            return {"signedURL": "https://storage.example/signed"}
    class Storage:
        def from_(self, bucket):
            assert bucket == "private-artifacts"
            return Bucket()
    class Client:
        storage = Storage()
    repository = SupabaseTaskRepository("", "", client=Client())
    repository.upload_artifact("private-artifacts", "task/file.xlsx", b"binary", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert repository.signed_artifact_url("private-artifacts", "task/file.xlsx") == "https://storage.example/signed"


class ArtifactRepository:
    def __init__(self, task, sources, findings):
        self.task, self.sources, self.findings, self.artifacts = task, sources, findings, {}

    def get_task(self, task_id): return self.task if task_id == self.task["id"] else None
    def list_sources(self, _task_id): return self.sources
    def list_findings(self, _task_id): return self.findings
    def save_artifact(self, row): self.artifacts[row["id"]] = row; return row
    def list_artifacts(self, _task_id): return list(self.artifacts.values())
    def get_artifact(self, artifact_id): return self.artifacts.get(artifact_id)
    def append_event(self, *_args): return {"sequence_no": 1}


@pytest.mark.parametrize("artifact_type", ["spreadsheet", "report", "email"])
def test_artifact_api_generates_persists_lists_and_downloads_local_file(tmp_path, monkeypatch, artifact_type):
    task, sources, findings = evidence_fixture()
    repo = ArtifactRepository(task, sources, findings)
    settings = Settings(_env_file=None, artifact_storage_backend="local", artifact_dir=str(tmp_path))
    import app.main as main_module
    monkeypatch.setattr(main_module, "settings", Settings(_env_file=None, app_env="test"))
    app.dependency_overrides[get_repository] = lambda: repo
    app.dependency_overrides[get_settings] = lambda: settings
    try:
        with TestClient(app) as client:
            response = client.post(f"/api/research/{task['id']}/artifacts", json={"artifact_type": artifact_type})
            assert response.status_code == 201, response.text
            artifact = response.json()
            assert "storage_path" not in artifact
            assert client.get(f"/api/research/{task['id']}/artifacts").json()["artifacts"][0]["id"] == artifact["id"]
            downloaded = client.get(f"/api/artifacts/{artifact['id']}/download")
            assert downloaded.status_code == 200
            if artifact_type == "spreadsheet":
                book = load_workbook(__import__("io").BytesIO(downloaded.content), read_only=False)
                assert book.sheetnames == ["Results", "Sources & Evidence", "Research Summary"]
                assert book["Results"]["A2"].value == "Alpha Clinic"
                assert book["Sources & Evidence"]["E2"].hyperlink.target == "https://example.org/fees"
            elif artifact_type == "report":
                assert b"# Research report" in downloaded.content and b"COMPLETED_WITH_LIMITATIONS" in downloaded.content
            else:
                assert b"has not been sent" in downloaded.content
    finally:
        app.dependency_overrides.clear()
